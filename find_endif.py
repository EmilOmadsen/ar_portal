from docx import Document

doc = Document('static/contract_templates/50_50 template med placeholders.docx')

print("All ENDIF paragraphs:")
for i, p in enumerate(doc.paragraphs):
    if 'endif' in p.text.lower():
        print(f"Para {i}: {p.text}")
